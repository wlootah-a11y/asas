"""The file becomes numbered lines, with no model; files that hold no CV are refused."""

import pytest
from conftest import docx

from asas_cv_profile import Lines, Refused, read_cv
from asas_cv_profile import lines as lines_mod
from asas_cv_profile.headings import Headings, chunks, marked, numbered_items


class FakePdf:
    def __init__(self, text: Lines):
        self.text, self.read_count = text, 0

    def read(self, data: bytes) -> Lines:
        self.read_count += 1
        return self.text


def test_lines_are_numbered_to_the_width_of_the_last():
    text = Lines([f"line {i}" for i in range(1, 1002)], "text")
    assert text.numbered().splitlines()[0] == "L0001 line 1"
    assert text.part([2, 1001]) == "L0002 line 2\nL1001 line 1001"
    assert text.line(0) is None and text.line(1002) is None


def test_an_almost_empty_file_is_refused():
    with pytest.raises(Refused) as exc:
        read_cv(docx("Page 1"), "cv.docx")
    assert exc.value.code == "document_too_little_text"


def test_word_keeps_document_order_with_tables_and_its_headers_first(monkeypatch):
    import io

    import docx as python_docx

    monkeypatch.setattr(lines_mod, "MIN_LETTERS", 0)
    d = python_docx.Document()
    d.sections[0].header.paragraphs[0].text = "Sara Example, sara@example.com"
    d.add_paragraph("Experience")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Analyst", "2021-2024"
    table.cell(1, 0).text, table.cell(1, 1).text = "Intern", "2020"
    d.add_paragraph("Education")
    buf = io.BytesIO()
    d.save(buf)
    text = read_cv(buf.getvalue(), "cv.docx")
    assert text.reader == "asas-cv"
    assert text.lines == ["Sara Example, sara@example.com", "Experience", "Analyst | 2021-2024", "Intern | 2020", "Education"]


def test_a_pdf_goes_to_the_pdf_reader_and_is_refused_without_one(monkeypatch):
    monkeypatch.setattr(lines_mod, "page_count", lambda data: 2)
    with pytest.raises(Refused) as exc:
        read_cv(b"%PDF-1.7", "cv.pdf")
    assert exc.value.code == "reader_not_configured"
    reader = FakePdf(Lines(["Sara Example, Senior Data Analyst at Example Co since 2019"] * 3, "docintel", pages=2))
    assert read_cv(b"%PDF-1.7", "cv.pdf", pdf=reader).reader == "docintel" and reader.read_count == 1


def test_a_long_file_that_is_not_a_cv_is_turned_away_before_or_after_reading(monkeypatch):
    chapters = Lines([f"Chapter {i}. The history of irrigation in the region, a study of canals and wells." for i in range(12)],
                     "docintel", pages=12)
    reader = FakePdf(chapters)
    monkeypatch.setattr(lines_mod, "page_count", lambda data: 12)
    with pytest.raises(Refused) as exc:
        read_cv(b"%PDF-1.7", "report.pdf", pdf=reader)
    assert exc.value.code == "document_not_a_cv"
    monkeypatch.setattr(lines_mod, "page_count", lambda data: 41)
    with pytest.raises(Refused) as exc:
        read_cv(b"%PDF-1.7", "report.pdf", pdf=reader)
    assert exc.value.code == "document_too_many_pages" and reader.read_count == 1  # turned away unread


def test_an_unknown_file_type_keeps_asas_cvs_code():
    with pytest.raises(Refused) as exc:
        read_cv(b"GIF89a", "photo.gif")
    assert exc.value.code == "document_unsupported_type"


def test_layout_tables_become_one_row_a_line():
    out = lines_mod.layout_markdown("## Publications\n<table><tr><th>TITLE</th><th>YEAR</th></tr><tr><td>A study</td><td>2022</td></tr></table>\n<!-- PageBreak -->")
    assert [x for x in out.splitlines() if x.strip()] == ["## Publications", "| TITLE | YEAR |", "| --- | --- |", "| A study | 2022 |"]


def test_journal_headings_editor_roles_and_courses_taught():
    h = Headings()
    assert h.heading("## REFERRED JOURNALS (SCI Indexed)") == {"publications"}
    assert h.heading("# EDITOR AND REVIEWER OF JOURNALS") == set()  # ends the list, starts none
    assert h.heading("COURSES TAUGHT") == set()  # teaching, not courses attended
    assert h.heading("Courses") == {"courses"} and h.heading("Honours and Awards") == {"awards"}
    assert h.heading("Project Manager") is None  # plain type, not only section names
    assert h.heading("##### ADDITIONAL INFORMATION") == set()


def test_a_host_adds_heading_words_for_a_kind():
    assert Headings().heading("PORTFOLIO") is None
    assert Headings({"projects": ["portfolio"]}).heading("PORTFOLIO") == {"projects"}


def test_a_split_list_runs_past_a_stray_heading():
    h = Headings()
    text = ["# PUBLICATIONS", "[1] A.", "GET MY OWN PROFILE", "| A study | 2022 |", "# TEACHING EXPERIENCE", "Lecturer"]
    assert h.find(text)["publications"] == [1, 2]
    assert h.find(text, other_ends=False)["publications"] == [1, 2, 3, 4]


def test_a_long_section_is_cut_where_items_start():
    text = []
    for i in range(1, 31):
        text += [f"[{i}] Example A. Study {i}.", "Sci Rep (2022)."]
    parts = chunks(list(range(1, 61)), text, size=25)
    assert [len(p) for p in parts] == [26, 26, 8]
    assert all(text[p[0] - 1].startswith("[") for p in parts)
    assert numbered_items(list(range(1, 61)), text) == 30


def test_a_markdown_heading_is_marked_whatever_its_words():
    assert marked("## Grants won") and not marked("## 2019 Grants") and not marked("Grants won")


def test_a_file_whose_bytes_are_not_the_type_its_name_claims_is_refused():
    reader = FakePdf(Lines(["Sara Example"] * 3, "docintel"))
    for data, name in ((b"GIF89a" + b"\0" * 64, "cv.pdf"), (b"%PDF-1.7", "cv.docx"), (b"PK\x03\x04 not a zip", "cv.docx")):
        with pytest.raises(Refused) as exc:
            read_cv(data, name, pdf=reader)
        assert exc.value.code == "document_unsupported_type"
    assert reader.read_count == 0  # turned away unread
    with pytest.raises(Refused) as exc:
        read_cv(b"", "cv.pdf", pdf=reader)
    assert exc.value.code == "document_empty"


def test_a_word_file_that_unpacks_too_far_is_refused_before_reading(monkeypatch):
    import io
    import zipfile

    monkeypatch.setattr(lines_mod, "MIN_LETTERS", 0)
    cv = docx("Sara Example")
    read_cv(cv, "cv.docx")
    buf = io.BytesIO(cv)
    with zipfile.ZipFile(buf, "a", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/media/padding.bin", b"\0" * 2**21)  # 2 MB of zeros packs to a few KB
    monkeypatch.setattr(lines_mod, "MAX_UNPACKED", 2**20)
    with pytest.raises(Refused) as exc:
        read_cv(buf.getvalue(), "cv.docx")
    assert exc.value.code == "document_unreadable"
    monkeypatch.setattr(lines_mod, "MAX_UNPACKED", 2**30)
    monkeypatch.setattr(lines_mod, "MAX_PARTS", 5)
    with pytest.raises(Refused) as exc:
        read_cv(cv, "cv.docx")
    assert exc.value.code == "document_unreadable"


def test_document_intelligence_reads_capped_pages_then_deletes_the_cv_and_closes(monkeypatch):
    di = pytest.importorskip("azure.ai.documentintelligence")
    from types import SimpleNamespace

    from azure.core.exceptions import ServiceRequestError, ServiceResponseError

    from asas_cv_profile import LayoutReader

    calls, retries = [], {}

    class FakeClient:
        def __init__(self, endpoint, credential, **kwargs):
            calls.append(("open", type(credential).__name__))
            retries.update(kwargs)

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            calls.append(("close",))

        def begin_analyze_document(self, model, body, pages, output_content_format, polling_interval=30):
            calls.append(("analyze", model, pages, polling_interval))
            def result(timeout=None):  # azure-core's LROPoller.result takes a timeout
                if fail_read:
                    raise ServiceResponseError("busy")
                return SimpleNamespace(content="# Sara Example\n\nSenior Data Analyst", pages=[1, 2])

            return SimpleNamespace(result=result, done=lambda: True, details={"operation_id": "op-1"})

        def delete_analyze_result(self, model, result_id):
            calls.append(("delete", model, result_id))
            if fail_delete:
                raise ServiceRequestError("down")

    monkeypatch.setattr(di, "DocumentIntelligenceClient", FakeClient)
    fail_delete = fail_read = False
    text = LayoutReader("https://example.cognitiveservices.azure.com", "k", max_pages=40).read(b"%PDF-1.7")
    assert text.lines == ["# Sara Example", "Senior Data Analyst"] and text.pages == 2
    assert calls == [("open", "AzureKeyCredential"), ("analyze", "prebuilt-layout", "1-40", 2),
                     ("delete", "prebuilt-layout", "op-1"), ("close",)]
    assert retries["retry_total"] == retries["retry_status"] == 6  # 429 and 5xx too, not azure-core's default of three
    calls.clear()
    fail_delete = True  # a failed delete keeps the read: Azure deletes it after 24 hours anyway
    entra = object()
    assert LayoutReader("https://example.cognitiveservices.azure.com", credential=entra).read(b"%PDF-1.7").pages == 2
    assert calls[0] == ("open", "object") and calls[-1] == ("close",)
    with pytest.raises(ValueError):
        LayoutReader("https://example.cognitiveservices.azure.com")
    with pytest.raises(ValueError):
        LayoutReader("https://example.cognitiveservices.azure.com", "k", credential=entra)
    calls.clear()
    fail_delete, fail_read = False, True  # a failed read is deleted too, not kept for 24 hours
    with pytest.raises(Refused) as exc:
        LayoutReader("https://example.cognitiveservices.azure.com", "k").read(b"%PDF-1.7")
    assert exc.value.code == "reader_unavailable" and calls[-2:] == [("delete", "prebuilt-layout", "op-1"), ("close",)]


def test_page_header_text_stays_and_a_running_header_goes():
    """Document Intelligence puts what it takes for a page header in a comment: a role's
    employer and dates at the top of page 2 must stay, the name on every page must not."""
    from asas_cv_profile.lines import layout_markdown

    md = ('<!-- PageHeader="Sara Example, CV" -->\n# Experience\nLecturer\n<!-- PageBreak -->\n'
          '<!-- PageHeader="Sara Example, CV" -->\n<!-- PageHeader="Example Bank &amp; Co" -->\n'
          '<!-- PageHeader="2010 - 2019" -->\nLed audits\n<!-- PageNumber="2" -->')
    out = [line for line in layout_markdown(md).splitlines() if line.strip()]
    assert out == ["# Experience", "Lecturer", "Example Bank & Co", "2010 - 2019", "Led audits"]


def test_a_read_that_does_not_finish_in_time_is_refused_as_unavailable(monkeypatch):
    """A stuck Document Intelligence operation must not hold a CV for ever: past read_timeout
    the file is refused as reader_unavailable, and the result is still deleted."""
    import sys
    import types

    from asas_cv_profile.lines import LayoutReader, Refused

    deleted = []

    class Poller:
        details = {"operation_id": "op-1"}

        def result(self, timeout=None):
            assert timeout == 0.01
            return None

        def done(self):
            return False

    class Client:
        def __init__(self, *a, **k):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def begin_analyze_document(self, *a, **k):
            return Poller()

        def delete_analyze_result(self, model, op):
            deleted.append(op)

    di = types.ModuleType("azure.ai.documentintelligence")
    di.DocumentIntelligenceClient = Client
    models = types.ModuleType("azure.ai.documentintelligence.models")
    models.DocumentContentFormat = types.SimpleNamespace(MARKDOWN="markdown")
    monkeypatch.setitem(sys.modules, "azure.ai.documentintelligence", di)
    monkeypatch.setitem(sys.modules, "azure.ai.documentintelligence.models", models)
    with pytest.raises(Refused) as caught:
        LayoutReader("https://example.cognitiveservices.azure.com", "key", read_timeout=0.01).read(b"%PDF-1.4")
    assert caught.value.code == "reader_unavailable" and deleted == ["op-1"]
