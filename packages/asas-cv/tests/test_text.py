"""Text out of an attached document, and what is refused.

The fixtures are BUILT here rather than checked in: a .docx is a zip of XML
and a PDF is a byte format, so constructing them with the same libraries that
will read them keeps the test honest about the shape the parser meets.
"""

import io

import pytest

from asas_cv import (
    CODES,
    MAX_DOCUMENT_BYTES,
    ExtractionError,
    cut_on_a_boundary,
    extract_text,
    normalize,
    resolve_content_type,
)


def docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, cell in enumerate(row):
                table.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pdf_bytes(lines: list[str]) -> bytes:
    """A minimal one-page PDF with a real text layer, drawn with Tj operators."""
    content = "BT /F1 12 Tf 72 720 Td 14 TL\n"
    for line in lines:
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content += f"({escaped}) Tj T*\n"
    content += "ET"
    stream = content.encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n".encode()
    out += b"%%EOF\n"
    return bytes(out)


class TestWordDocuments:
    def test_reads_the_paragraphs(self) -> None:
        text = extract_text(docx_bytes(["Senior Data Engineer", "Owns the ingestion platform."]), filename="cv.docx")
        assert "Senior Data Engineer" in text
        assert "Owns the ingestion platform." in text

    def test_reads_the_tables_too(self) -> None:
        """python-docx keeps table text out of `paragraphs` entirely; a
        paragraphs-only read would drop a CV's experience table silently."""
        data = docx_bytes(["Experience"], table_rows=[["Role", "Detail"], ["Engineer", "Owns the Kafka pipelines"]])
        assert "Owns the Kafka pipelines" in extract_text(data, filename="cv.docx")

    def test_keeps_the_document_order_so_a_table_stays_under_its_heading(self) -> None:
        data = docx_bytes(["1. Profile", "2. Experience", "3. Notes"], table_rows=[["Grade", "G4"]])
        lines = [line for line in extract_text(data, filename="cv.docx").split("\n") if line.strip()]
        assert lines.index("1. Profile") < lines.index("Grade | G4")
        assert lines.index("2. Experience") < lines.index("Grade | G4")

    def test_falls_back_rather_than_failing_if_python_docx_internals_move(self, monkeypatch) -> None:
        import asas_cv.text as mod

        def boom(_document):
            raise AttributeError("internals moved")
            yield  # pragma: no cover

        monkeypatch.setattr(mod, "_docx_blocks", boom)
        text = extract_text(docx_bytes(["Platform Engineer"], table_rows=[["Grade", "G4"]]), filename="cv.docx")
        assert "Platform Engineer" in text and "Grade | G4" in text

    def test_trusts_the_extension_over_the_browsers_content_type(self) -> None:
        text = extract_text(docx_bytes(["Platform Engineer"]), filename="cv.docx", declared_type="application/octet-stream")
        assert "Platform Engineer" in text


class TestPdfDocuments:
    def test_reads_a_text_layer(self) -> None:
        assert "Head of Data Platform" in extract_text(pdf_bytes(["Head of Data Platform", "Reports to the CTO."]), filename="cv.pdf")

    def test_a_scan_is_refused_by_name_not_as_a_generic_failure(self) -> None:
        with pytest.raises(ExtractionError) as caught:
            extract_text(pdf_bytes([]), filename="scan.pdf")  # valid, no text drawn
        assert caught.value.code == "document_no_text"

    def test_a_damaged_pdf_says_unreadable_rather_than_raising_pdfminers_own(self) -> None:
        with pytest.raises(ExtractionError) as caught:
            extract_text(b"%PDF-1.4\nnot really a pdf", filename="broken.pdf")
        assert caught.value.code in {"document_unreadable", "document_no_text"}


class TestPlainText:
    def test_reads_utf8(self) -> None:
        assert "Café" in extract_text("Café lead".encode(), filename="cv.txt")

    def test_reads_arabic_saved_from_windows_notepad(self) -> None:
        arabic = "مهندس بيانات أول"
        assert extract_text(arabic.encode("cp1256"), filename="cv.txt") == arabic

    def test_utf16_only_on_a_bom(self) -> None:
        assert extract_text("Résumé".encode("utf-16"), filename="cv.txt") == "Résumé"

    def test_markdown_is_plain_text(self) -> None:
        assert "Owns ingestion" in extract_text(b"# Role\n\n- Owns ingestion\n", filename="cv.md")


class TestRefusals:
    def test_an_unsupported_type_is_named(self) -> None:
        with pytest.raises(ExtractionError) as caught:
            resolve_content_type("photo.png", "image/png")
        assert caught.value.code == "document_unsupported_type"

    def test_a_declared_type_is_consulted_only_without_an_extension(self) -> None:
        assert resolve_content_type("upload", "application/pdf") == "application/pdf"
        assert resolve_content_type("cv.pdf", "image/png") == "application/pdf"

    def test_an_empty_file(self) -> None:
        with pytest.raises(ExtractionError) as caught:
            extract_text(b"", filename="cv.pdf")
        assert caught.value.code == "document_empty"

    def test_a_file_over_the_cap(self) -> None:
        with pytest.raises(ExtractionError) as caught:
            extract_text(b"x" * (MAX_DOCUMENT_BYTES + 1), filename="cv.txt")
        assert caught.value.code == "document_too_large"

    def test_the_cap_is_the_callers_to_set(self) -> None:
        with pytest.raises(ExtractionError) as caught:
            extract_text(b"x" * 11, filename="cv.txt", max_bytes=10)
        assert caught.value.code == "document_too_large"

    def test_the_size_cap_is_checked_before_the_parse(self) -> None:
        with pytest.raises(ExtractionError) as caught:
            extract_text(b"%PDF" + b"x" * MAX_DOCUMENT_BYTES, filename="huge.pdf")
        assert caught.value.code == "document_too_large"

    def test_every_code_is_declared(self) -> None:
        assert ExtractionError("x", code="document_empty").code in CODES
        with pytest.raises(ValueError):
            ExtractionError("x", code="made_up")


class TestCutting:
    def test_a_short_text_passes_whole(self) -> None:
        assert cut_on_a_boundary("Owns the ingestion platform.", 100) == "Owns the ingestion platform."

    def test_the_cut_falls_on_a_paragraph_break_not_mid_word(self) -> None:
        body = ("A" * 500 + "\n\n") * 200
        text = cut_on_a_boundary(body, 10_000)
        assert text.endswith("A") and len(text) <= 10_000

    def test_a_document_with_no_breaks_at_all_still_yields_the_budget(self) -> None:
        assert len(cut_on_a_boundary("A" * 20_000, 10_000)) == 10_000


def test_normalization_tidies_whitespace_without_joining_paragraphs() -> None:
    assert normalize("Heading\r\n\r\n\r\n\r\nBody line   \r\nMore\n") == "Heading\n\nBody line\nMore"
    assert extract_text(b"Heading\r\n\r\n\r\n\r\nBody line   \r\nMore\n", filename="s.txt") == "Heading\n\nBody line\nMore"
