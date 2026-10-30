"""Bytes out of an attached document: PDF, Word (.docx), or plain text.

Mechanical, bytes in and characters out, not an AI act. What it REFUSES
matters as much as what it reads, because the result is usually handed to a
model or a search index: each refusal is an :class:`ExtractionError` with a
stable ``code`` a host maps to its own wording.

What a real upload surface taught this module:

- **Trust the extension, not the browser's content type.** Windows sends
  ``application/octet-stream`` for a .docx often enough that trusting it
  would refuse the single most common document format; the declared type
  is consulted only when the name carries no useful extension.
- **A Word document's substance sits in TABLES**, which python-docx keeps out
  of ``paragraphs`` entirely, and ``paragraphs`` and ``tables`` are two
  separate sequences, so reading one then the other moves every table to the
  end of the file. The body XML is walked in DOCUMENT order instead.
- **A scan is a real document with no text in it**, and the remedy (type the
  text) is not the remedy for a corrupt upload, so it is its own code.
- **The size cap is checked before the parse**, or an eleven-megabyte PDF is
  fully parsed and only then refused.
- **Text files: UTF-16 only on a BOM**, because without one it decodes almost
  any even-length byte string into garbage instead of raising; cp1256 before
  latin-1, because Arabic text saved from Windows Notepad is a realistic input
  and latin-1 decodes ANY byte, so it must be last.
"""

from __future__ import annotations

import codecs
import io
import re
import unicodedata

#: Ten megabytes: generous for a document a person wrote by hand. One cap, so
#: every upload surface in a product refuses at one size.
MAX_DOCUMENT_BYTES = 10 * 1024 * 1024

PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_CONTENT_TYPES = frozenset({"text/plain", "text/markdown", "text/x-markdown"})
_EXTENSIONS = {
    ".pdf": PDF_CONTENT_TYPE,
    ".docx": DOCX_CONTENT_TYPE,
    ".txt": "text/plain",
    ".md": "text/markdown",
}

# Collapse the run of blank lines a PDF extractor leaves between blocks,
# without joining paragraphs: structure (headings, bullet lists) is signal a
# caller may still want, so this tidies whitespace and nothing else.
_BLANK_RUN = re.compile(r"\n{3,}")
_TRAILING_WS = re.compile(r"[ \t]+\n")

CODES = (
    "document_unsupported_type",
    "document_empty",
    "document_too_large",
    "document_unreadable",
    "document_no_text",
)
"""Every ``ExtractionError.code``."""


class ExtractionError(ValueError):
    """The file cannot become text. ``code`` is one of :data:`CODES`; a host
    with its own error vocabulary catches this and re-raises its own."""

    def __init__(self, message: str, *, code: str) -> None:
        if code not in CODES:
            raise ValueError(f"unknown extraction code {code!r}")
        super().__init__(message)
        self.code = code


def resolve_content_type(filename: str, declared: str | None) -> str:
    """The type the document will be PARSED as, from its extension first."""
    lowered = (filename or "").lower()
    for ext, content_type in _EXTENSIONS.items():
        if lowered.endswith(ext):
            return content_type
    if declared == PDF_CONTENT_TYPE or declared == DOCX_CONTENT_TYPE or declared in TEXT_CONTENT_TYPES:
        return declared
    raise ExtractionError(
        "This must be a PDF, a Word (.docx) document, or plain text.",
        code="document_unsupported_type",
    )


def extract_text(
    data: bytes,
    *,
    filename: str,
    declared_type: str | None = None,
    max_bytes: int = MAX_DOCUMENT_BYTES,
) -> str:
    """Bytes to the document's text, normalised. Raises :class:`ExtractionError`."""
    if not data:
        raise ExtractionError("The file is empty.", code="document_empty")
    if len(data) > max_bytes:
        raise ExtractionError(
            f"The file is larger than {max_bytes // (1024 * 1024)} MB.",
            code="document_too_large",
        )

    content_type = resolve_content_type(filename, declared_type)
    if content_type == PDF_CONTENT_TYPE:
        text = _from_pdf(data)
    elif content_type == DOCX_CONTENT_TYPE:
        text = _from_docx(data)
    else:
        text = _from_plain(data)

    text = normalize(text)
    if not text:
        raise ExtractionError(
            "No readable text could be found in this document. If it is a "
            "scan, type the text instead.",
            code="document_no_text",
        )
    return text


def cut_on_a_boundary(text: str, limit: int) -> str:
    """``text`` shortened to ``limit``, ending on a paragraph or line break
    where one falls in the last quarter of the window, so nothing downstream
    reads half a sentence as a whole one."""
    window = text[:limit]
    cut = window.rfind("\n\n")
    if cut < limit * 3 // 4:
        cut = window.rfind("\n")
    if cut < limit * 3 // 4:
        cut = limit
    return window[:cut].rstrip()


def normalize(text: str) -> str:
    """NFC, Unix line ends, no NULs, no trailing spaces, blank runs collapsed
    to one blank line, trimmed."""
    text = unicodedata.normalize("NFC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = _TRAILING_WS.sub("\n", text)
    text = _BLANK_RUN.sub("\n\n", text)
    return text.strip()


def _from_pdf(data: bytes) -> str:
    from pdfminer.high_level import extract_text as pdf_extract  # local: slow import

    try:
        return pdf_extract(io.BytesIO(data)) or ""
    except Exception as exc:  # pdfminer raises a wide family on a damaged file
        raise ExtractionError("This PDF could not be read.", code="document_unreadable") from exc


def _docx_blocks(document: object):
    """Paragraphs and tables in DOCUMENT ORDER, off the body XML (python-docx
    has no public in-order iterator)."""
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body  # type: ignore[attr-defined]
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _from_docx(data: bytes) -> str:
    import docx  # local: slow import
    from docx.table import Table

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(
            "This Word document could not be read.", code="document_unreadable"
        ) from exc

    def render_table(table: Table) -> list[str]:
        # Rows are pipe-joined, which reads as a row to a model and needs no
        # layout.
        out = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(" | ".join(cells))
        return out

    blocks: list[str] = []
    try:
        for item in _docx_blocks(document):
            if isinstance(item, Table):
                blocks.extend(render_table(item))
            else:
                blocks.append(item.text)
    except Exception:  # noqa: BLE001 - internals moved; a scrambled document still beats none
        blocks = [p.text for p in document.paragraphs]
        for table in document.tables:
            blocks.extend(render_table(table))
    return "\n".join(blocks)


def _from_plain(data: bytes) -> str:
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16")
    if data.startswith(codecs.BOM_UTF8):
        return data.decode("utf-8-sig")
    for encoding in ("utf-8", "cp1256", "latin-1"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    raise ExtractionError(
        "This text file's encoding could not be read.", code="document_unreadable"
    )
